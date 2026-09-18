"""视图层集成测试：geojson / 导出 / 查询分页过滤 / 统计API / 统计页。

这些视图都是 @login_required（非管理员专属），用普通用户即可访问。
"""
import io
import re

import openpyxl
import pytest


@pytest.mark.django_db
def test_geojson_point_has_gps_ts(client, regular_user, track_factory):
    """回放时长换算依赖每个点的 gps_ts，这里锁定该字段存在。"""
    t = track_factory(points=[{"x": 0, "y": 0}, {"x": 10, "y": 0}, {"x": 20, "y": 0}])
    client.force_login(regular_user)
    resp = client.get(f"/tracks/geojson/{t.pk}/")
    assert resp.status_code == 200
    data = resp.json()
    types = [f["geometry"]["type"] for f in data["features"]]
    assert "LineString" in types
    pts = [f for f in data["features"] if f["geometry"]["type"] == "Point"]
    assert len(pts) == 3
    assert all("gps_ts" in f["properties"] for f in pts)


@pytest.mark.django_db
def test_export_statistics_returns_valid_xlsx(client, regular_user, track_factory):
    track_factory(points=[{"x": 0, "y": 0}, {"x": 10, "y": 0}])
    client.force_login(regular_user)
    resp = client.get("/tracks/statistics/export/")
    assert resp.status_code == 200
    assert "spreadsheet" in resp["Content-Type"]
    # 文件名是中文，Django 对 Content-Disposition 做 RFC 2047 编码，需先解码再断言
    from email.header import decode_header, make_header
    disposition = str(make_header(decode_header(resp["Content-Disposition"])))
    assert "attachment" in disposition
    assert ".xlsx" in disposition
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws.cell(1, 1).value == "地块名称"
    assert ws.max_row >= 2  # 表头 + 至少一行数据


@pytest.mark.django_db
def test_track_query_pagination(client, regular_user, track_factory):
    # 25 个点，每页 20 → 应 2 页
    t = track_factory(points=[{"x": i, "y": 0} for i in range(25)])
    client.force_login(regular_user)
    resp = client.get(f"/tracks/query/?track_id={t.pk}")
    assert resp.status_code == 200
    body = resp.content.decode()
    assert "轨迹点数据" in body
    assert "共 25 条" in body
    assert "1/2" in body  # 第 1 页 / 共 2 页


@pytest.mark.django_db
def test_track_query_filter_by_working_status(client, regular_user, track_factory):
    t = track_factory(points=[
        {"x": 0, "y": 0, "working": True},
        {"x": 1, "y": 0, "working": True},
        {"x": 2, "y": 0, "working": False},
    ])
    client.force_login(regular_user)

    def filtered_count(body):
        m = re.search(r"筛选后 (\d+) 条", body)
        return int(m.group(1)) if m else None

    all_body = client.get(f"/tracks/query/?track_id={t.pk}").content.decode()
    work_body = client.get(f"/tracks/query/?track_id={t.pk}&status=working").content.decode()
    assert filtered_count(all_body) == 3
    assert filtered_count(work_body) == 2


@pytest.mark.django_db
def test_track_stats_api_json_shape(client, regular_user, track_factory):
    t = track_factory(points=[{"x": 0, "y": 0}, {"x": 10, "y": 0}], width=2)
    client.force_login(regular_user)
    resp = client.get(f"/tracks/stats/{t.pk}/")
    assert resp.status_code == 200
    s = resp.json()
    for key in ["name", "total_duration_hours", "total_distance_km", "area_mu",
                "avg_working_speed", "point_count", "working_point_count",
                "shutdown_count", "width"]:
        assert key in s
    # 面积 = 10m × 2 = 20 m² → 亩
    assert s["area_mu"] == pytest.approx(20 / 666.667, abs=0.01)


@pytest.mark.django_db
def test_statistics_page_renders_summary(client, regular_user, track_factory):
    track_factory(points=[{"x": 0, "y": 0}, {"x": 10, "y": 0}])
    client.force_login(regular_user)
    resp = client.get("/tracks/statistics/")
    assert resp.status_code == 200
    body = resp.content.decode()
    assert "作业总时长" in body
    assert "导出Excel" in body
    assert "清洗前后" in body  # 清洗对比图标题
